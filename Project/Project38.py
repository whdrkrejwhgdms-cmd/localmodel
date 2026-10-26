import requests
from bs4 import BeautifulSoup
import openpyxl
from datetime import datetime

# 네이버 환율 정보 웹 페이지 URL
url = "https://search.naver.com/search.naver?&query=환율"

# 웹 페이지 요청 및 HTML 파싱
response = requests.get(url)
soup = BeautifulSoup(response.text, 'html.parser')

# 크롤링 대상 국가 목록
currency_list = ['미국(USD)', '일본(JPY)', '유럽연합(EUR)',
                 '중국(CNY)', '영국(GBP)', '호주(AUD)']

# 새 Excel 워크북 생성 및 시트 활성화
workbook = openpyxl.Workbook()
sheet = workbook.active

# 파일 이름 및 시트 이름을 위한 오늘 날짜 가져오기
today_date = datetime.now().strftime("%Y%m%d")

# 시트에 헤더 작성
sheet.cell(row=1, column=1, value="국가")
sheet.cell(row=1, column=2, value="환율")
sheet.cell(row=1, column=3, value="전일 대비")
sheet.cell(row=1, column=4, value="증감율")

# 각 국가에 대한 환율 정보를 시트에 작성
exchange_rate = soup.find_all('span', {'class': 'spt_con'})
for i in range(min(len(currency_list), len(exchange_rate))):
    currency = currency_list[i]
    rate_info = exchange_rate[i].text.split(' ')

    # 시트에 데이터 작성
    sheet.cell(row=i+2, column=1, value=currency)
    sheet.cell(row=i+2, column=2, value=rate_info[1])
    sheet.cell(row=i+2, column=3, value=rate_info[3])
    sheet.cell(row=i+2, column=4, value=rate_info[4])

# 시트 이름을 오늘 날짜로 설정
sheet.title = today_date

# 오늘 날짜가 포함된 파일 경로에 워크북 저장
file_path = f'C:\\AutoCoder\\exchange_rates_{today_date}.xlsx'
workbook.save(file_path)
print(f"엑셀 파일이 {file_path}에 저장되었습니다.")