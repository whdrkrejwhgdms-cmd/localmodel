from webdriver_manager.chrome import ChromeDriverManager
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
import time
import os

# 크롬 옵션 지정 (자동 다운로드)
chrome_options = webdriver.ChromeOptions()
chrome_options.add_argument('--enable-download')
driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()),
                          options=chrome_options)

# C社의 식품 카테고리로 이동
driver.get(url="https://www.coupang.com/np/categories/194276")
time.sleep(5)

# Excel 파일 생성
from openpyxl import Workbook
wb = Workbook()
ws = wb.create_sheet('식품')
del (wb['Sheet'])
ws.append(['이름', '가격', '배송기한', 'URL'])

# 'productList' ID를 가진 요소에서 모든 li 태그 찾기
lis = driver.find_element(By.ID, 'productList').find_elements(By.TAG_NAME, 'li')

# 각 li 태그를 순회하며 정보 추출
for li in lis:
    try:
        # 제품명 추출
        product = li.find_element(By.CLASS_NAME, 'name').text
        # 가격 추출
        price = li.find_element(By.CLASS_NAME, 'price-value').text
        # 배송 정보 추출
        delivery = li.find_element(By.CLASS_NAME, 'delivery').text
        # 제품 URL 추출
        url = li.find_element(By.CLASS_NAME, 'baby-product-link').get_attribute('href')
        # 추출한 정보를 Excel 시트에 추가
        ws.append([product, price, delivery, url])

    # 예외 발생 시 무시
    except Exception:
        pass

# Excel 파일 저장
wb.save('C:/AutoCoder/Coupang_Excel.xlsx')
# Excel 파일 닫기
wb.close()

# 기존 Excel 파일 로딩
from openpyxl import load_workbook
WB = load_workbook('C:/AutoCoder/Coupang_Excel.xlsx')
# 활성 시트 가져오기
WS = WB.active

# 2번째 행부터 최대 행까지 반복
for i in range(2, WS.max_row + 1):
    # A열에 하이퍼링크 삽입
    WS['A' + str(i)].hyperlink = WS['D' + str(i)].value
    # 하이퍼링크 스타일 설정
    WS['A' + str(i)].style = "Hyperlink"

# D열 삭제
WS.delete_cols(4)
# 하이퍼링크가 적용된 새 파일 저장
WB.save('C:/AutoCoder/Coupang_Excel_HyperLink.xlsx')
# 원본 파일 닫기
WB.close()
# 원본 파일 삭제
os.remove('C:/AutoCoder/Coupang_Excel.xlsx')