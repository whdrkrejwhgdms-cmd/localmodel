# 필요한 라이브러리 임포트
from os import listdir
from os.path import join
from openpyxl import load_workbook, Workbook

# 작업 디렉터리 경로 지정
directory_path = "C:\\AutoCoder"
files = listdir(directory_path)

# 결과를 저장할 Excel 파일을 만듭니다.
result_xlsx = Workbook()
result_sheet = result_xlsx.active

# 각 파일을 순서대로 처리합니다.
for myfile in files:
    # 파일 확장자가 xlsx가 아닌 경우 건너뜁니다.
    if myfile[-4:] != "xlsx":
        continue

    # 파일 경로를 생성합니다.
    file_path = join(directory_path, myfile)

    # 대상 파일을 읽기 전용으로 열어봅니다.
    tg_xlsx = load_workbook(file_path, read_only=True)
    tg_sheet = tg_xlsx.active

    # 대상 파일의 각 행을 순서대로 처리합니다.
    for row in tg_sheet.iter_rows():
        # 행 데이터를 담을 배열을 생성합니다.
        row_data = []

        # 각 행의 각 셀을 순서대로 처리합니다.
        for cell in row:
            # 셀의 값을 배열에 추가합니다.
            row_data.append(cell.value)

        # 결과 파일의 한 행에 대상 파일의 한 행의 데이터를 추가합니다.
        result_sheet.append(row_data)

# 결과 파일을 저장합니다.
result_xlsx.save(join(directory_path, "mergedLog.xlsx"))